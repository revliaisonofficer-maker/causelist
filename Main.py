import os
import sys
import time
import re
import pandas as pd
import copy
from selenium import webdriver
from selenium.webdriver.firefox.service import Service
from selenium.webdriver.firefox.options import Options
from webdriver_manager.firefox import GeckoDriverManager
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import UnexpectedAlertPresentException, NoAlertPresentException, TimeoutException

# --- OPENPYXL ---
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# Force standard output to UTF-8 to prevent Windows terminal character map crashes with emojis
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# =========================================================
# ⚙️ SETTINGS - CONTROL YOUR RUN HERE
# =========================================================
SELECTED_COURTS = [] 
TARGET_DISTRICTS = ["SRIKAKULAM", "VIZIANAGARAM"] 
# =========================================================

# --- CONFIG ---
OUTPUT_DIR = os.path.normpath(r'D:\Causelist Automation')
os.makedirs(OUTPUT_DIR, exist_ok=True)
FILE_2_PATH = r"D:\Causelist Automation\Complete List for Vlookup.xlsx"
F2_SHEET_1, F2_SHEET_2 = '1', '2'
SHEET_NAME = 'Case and Advocates'
COMPARE_COLUMN = 'Case No'

YELLOW_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
BOLD_FONT = Font(bold=True)

CASE_NO_PATTERN = re.compile(r"([A-Z()]+\s*/\s*\d{1,5}\s*/\s*\d{4})")
S_NO_PATTERN = re.compile(r"^\s*(\d{1,3})")

def _clean(text):
    return " ".join(str(text).split()).strip()

def main():
    print(f"🚀 Initializing Headless Engine (Invisible Mode)...")
    
    options = Options()
    # --- HEADLESS MODE ENABLED WITH BROWSER MASQUERADING ---
    options.add_argument("-headless") 
    options.add_argument("--window-size=1920,1080")
    options.set_preference("general.useragent.override", "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:123.0) Gecko/20100101 Firefox/123.0")
    
    driver = webdriver.Firefox(service=Service(GeckoDriverManager().install()), options=options)
    wait = WebDriverWait(driver, 25)

    try:
        driver.get("https://aphc.gov.in/Hcdbs/search.jsp")
        time.sleep(3)
        
        # Robust Text-Based Selection instead of fragile structure paths (:nth-child)
        print("🔗 Navigating to Daily Cause List...")
        daily_btn = wait.until(EC.element_to_be_clickable(
            (By.XPATH, "//button[contains(normalize-space(), 'Daily') or contains(normalize-space(), 'Daily Cause List')]")
        ))
        daily_btn.click()
        time.sleep(2)
        
        # Click the submit/search button
        wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, ".btn-gradient-info"))).click()
        
        try:
            time.sleep(2)
            driver.switch_to.alert.accept()
            print("⚠️ No lists found for today.")
            return 
        except NoAlertPresentException:
            pass

        dropdown_el = wait.until(EC.presence_of_element_located((By.ID, "court")))
        all_options = Select(dropdown_el).options
        
        court_map = {}
        for opt in all_options:
            val, txt = opt.get_attribute('value'), opt.text
            if val == '0': continue
            if SELECTED_COURTS:
                num_match = re.search(r'\d+', txt)
                if num_match and num_match.group() in SELECTED_COURTS:
                    court_map[val] = txt
            else:
                court_map[val] = txt

        if not court_map:
            print("❌ No matching courts found!")
            return

        print(f"📡 Processing {len(court_map)} courts in background...")
        
        all_data = []
        for val, name in court_map.items():
            try:
                print(f"⚖️ {name}...", end=" ", flush=True)
                dropdown_el = wait.until(EC.presence_of_element_located((By.ID, "court")))
                Select(dropdown_el).select_by_value(val)
                time.sleep(4) 
                
                rows = driver.find_elements(By.CSS_SELECTOR, "table tr")
                current_cat = ""
                for row in rows:
                    ths = row.find_elements(By.TAG_NAME, "th")
                    if ths:
                        current_cat = _clean(ths[0].text)
                        continue
                    tds = row.find_elements(By.TAG_NAME, "td")
                    if len(tds) < 6: continue
                    
                    s_no = _clean(tds[0].text)
                    if S_NO_PATTERN.match(s_no):
                        case_match = CASE_NO_PATTERN.search(_clean(tds[1].text))
                        case_no = re.sub(r'\s+', '', case_match.group(1)) if case_match else "N/A"
                        all_data.append({
                            'Category': current_cat, 'Court': name, 'S.No': s_no, 
                            'Case No': case_no, 'Party Name': _clean(tds[2].text),
                            'Petitioner Advocate': _clean(tds[3].text),
                            'Respondent Advocate': _clean(tds[4].text), 'District': _clean(tds[5].text).upper()
                        })
                driver.back()
                time.sleep(2)
            except Exception:
                continue

        if all_data:
            output_file = os.path.join(OUTPUT_DIR, f"Causelist_{time.strftime('%Y%m%d_%H%M')}.xlsx")
            df_live = pd.DataFrame(all_data)
            df_live['Clean_Case'] = df_live[COMPARE_COLUMN].astype(str).str.replace(r'\s+', '', regex=True)

            master_y = pd.read_excel(FILE_2_PATH, sheet_name=F2_SHEET_1)
            master_r = pd.read_excel(FILE_2_PATH, sheet_name=F2_SHEET_2)
            master_y['Clean_Case'] = master_y[COMPARE_COLUMN].astype(str).str.replace(r'\s+', '', regex=True)
            master_r['Clean_Case'] = master_r[COMPARE_COLUMN].astype(str).str.replace(r'\s+', '', regex=True)

            # 1. Main Data Sheet
            df_live.drop(columns=['Clean_Case']).to_excel(output_file, sheet_name=SHEET_NAME, index=False)
            wb = load_workbook(output_file)
            ws_main = wb[SHEET_NAME]
            
            y_clean_set = set(master_y['Clean_Case'].tolist())
            r_clean_set = set(master_r['Clean_Case'].tolist())

            for row in range(2, ws_main.max_row + 1):
                case_val = str(ws_main.cell(row=row, column=4).value).replace(" ", "")
                if case_val in y_clean_set: ws_main.cell(row=row, column=4).fill = YELLOW_FILL
                if case_val in r_clean_set: ws_main.cell(row=row, column=4).fill = RED_FILL

            # 2. Filtered District and Master sheet
            print("\n📂 Creating Filtered District and Master sheet...")
            ws_filter = wb.create_sheet("Filtered District and Master")
            for col in range(1, ws_main.max_column + 1):
                ws_filter.cell(row=1, column=col, value=ws_main.cell(row=1, column=col).value).font = BOLD_FONT

            target_row = 2
            for row in range(2, ws_main.max_row + 1):
                case_val = str(ws_main.cell(row=row, column=4).value).replace(" ", "")
                dist_val = str(ws_main.cell(row=row, column=8).value or "").upper()
                if (case_val in y_clean_set or case_val in r_clean_set) or any(d in dist_val for d in TARGET_DISTRICTS):
                    for col in range(1, ws_main.max_column + 1):
                        new_cell = ws_filter.cell(row=target_row, column=col, value=ws_main.cell(row=row, column=col).value)
                        if ws_main.cell(row=row, column=col).has_style:
                            new_cell.fill = copy.copy(ws_main.cell(row=row, column=col).fill)
                    target_row += 1

            # 3 & 4. Mirror Verification Sheets
            def add_full_mirror_sheet(label, master_df):
                v_ws = wb.create_sheet(label)
                merged = pd.merge(master_df, df_live, on='Clean_Case', suffixes=('_Master', '_Live'))
                if merged.empty: return
                m_keys = [c for c in merged.columns if c.endswith('_Master') or (c in master_df.columns and c != 'Clean_Case' and f"{c}_Master" not in merged.columns)]
                l_keys = [c for c in merged.columns if c.endswith('_Live') or (c in df_live.columns and c != 'Clean_Case' and f"{c}_Live" not in merged.columns)]
                all_headers = [k.replace('_Master', '') for k in m_keys] + [" | "] + [k.replace('_Live', '') for k in l_keys]
                for c_idx, h_text in enumerate(all_headers, 1):
                    v_ws.cell(row=1, column=c_idx, value=h_text).font = BOLD_FONT
                for r_idx, row_dict in enumerate(merged.to_dict('records'), 2):
                    for c_idx, key in enumerate(m_keys, 1):
                        v_ws.cell(row=r_idx, column=c_idx, value=row_dict[key])
                    v_ws.cell(row=r_idx, column=len(m_keys) + 1, value="||")
                    for c_idx, key in enumerate(l_keys, 1):
                        v_ws.cell(row=r_idx, column=len(m_keys) + 2 + c_idx - 1, value=row_dict[key])

            add_full_mirror_sheet("Master Yellow Verification", master_y)
            add_full_mirror_sheet("Master Red Verification", master_r)
            
            wb.save(output_file)
            print(f"\n✅ COMPLETE. 4 sheets saved to {output_file}")
            
    except TimeoutException:
        print("\n❌ Element search timed out. Taking an error snapshot...")
        screenshot_path = os.path.join(OUTPUT_DIR, "selenium_timeout_error.png")
        driver.save_screenshot(screenshot_path)
        print(f"📸 Snapshot dropped in '{screenshot_path}'. Open it to check the website state.")
        raise

    finally:
        driver.quit()

if __name__ == "__main__":
    main()
